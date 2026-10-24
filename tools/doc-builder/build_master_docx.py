#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
MoBo Dual-Site Master Manual DOCX Builder (v2.1)
路径: /opt/mobo-stack/scripts/build_master_docx.py
功能: 无需 Pandoc，纯 Python (python-docx) 将双站系统总册 Markdown 或 /opt/mobo-stack 实盘源码
      直接排版生成机构级 Word 文档《莫伯双站系统全栈架构与部署总册_v2.1.docx》
依赖: pip install python-docx
"""

import os
import re
import sys
import argparse
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================================
# 1. 华尔街机构级配色常量 (深海蓝 + 香槟金 + 柔和灰)
# ============================================================================
COLOR_NAVY_DARK = RGBColor(13, 24, 46)       # #0D182E 一级/二级标题与表头深蓝
COLOR_NAVY_MED = RGBColor(30, 51, 92)        # #1E335C 三级标题蓝
COLOR_GOLD_ACCENT = RGBColor(197, 160, 89)   # #C5A059 莫伯资本香槟金
COLOR_TEXT_MAIN = RGBColor(30, 41, 59)       # #1E293B 正文深石板灰
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)   # #64748B 页眉页脚与注释灰
COLOR_CODE_TEXT = RGBColor(15, 23, 42)       # #0F172A 代码深色字体

HEX_NAVY_HEADER = "0D182E"                   # 表头背景深蓝
HEX_ZEBRA_ROW = "F8FAFC"                     # 表格交替斑马纹浅灰
HEX_CODE_BG = "F1F5F9"                       # 代码块容器背景色
HEX_QUOTE_BG = "FEFCE8"                      # 引用提示框暖黄底色
HEX_BORDER_GRAY = "CBD5E1"                   # 表格与容器边框灰
HEX_GOLD_LINE = "C5A059"                     # 标题底部金线

# --- 封面/页眉文案（可用命令行覆盖，使本工具可渲染任意 Markdown 文档）-----
DOC_TAG = "ARCHITECTURE & FULL-STACK ENGINEERING MANUAL · V2.1"
DOC_TITLE = "莫伯双站系统全栈架构与部署总册"
DOC_SUBTITLE = (
    "MoBoCapital.com（莫伯资本投研智库与对冲基金孵化平台）\n"
    "MoBoHome.com（莫伯之家个人主权生命数字档案馆）"
)
HEADER_TEXT = "MO BO CAPITAL & MO BO HOME  |  莫伯双站系统全栈架构与部署总册 (v2.1)"
COVER_META_ROWS: list[tuple[str, str]] = []   # 非空时覆盖默认封面元数据卡


# ============================================================================
# 2. 底层 OXML 样式与中文字体绑定工具函数
# ============================================================================
def set_run_font(run, ascii_font="Inter", east_asia_font="Microsoft YaHei", size_pt=10.5, bold=False, italic=False, color=COLOR_TEXT_MAIN):
    """强制绑定中西文字体，彻底解决 python-docx 中文回退到默认宋体的问题"""
    run.font.name = ascii_font
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:cs"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east_asia_font)


def set_cell_background(cell, hex_color: str):
    """设置 Word 表格单元格背景填充色"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """设置单元格内边距 (单位: dxa, 20 dxa = 1 pt)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color_hex=HEX_BORDER_GRAY, sz="4"):
    """为表格设置柔和的细实线边框"""
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for b_name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement(f"w:{b_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), sz)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color_hex)
        borders.append(border)
    tblPr.append(borders)


def add_bottom_border_to_paragraph(paragraph, color_hex=HEX_GOLD_LINE, sz="12"):
    """为一级/二级标题底部添加机构风装饰线"""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)


def enable_auto_update_fields(doc):
    """开启打开 Word 时自动更新目录 (TOC) 与页码域"""
    namespace = "{[http://schemas.openxmlformats.org/wordprocessingml/2006/main](http://schemas.openxmlformats.org/wordprocessingml/2006/main)}"
    element_update_fields = OxmlElement("w:updateFields")
    element_update_fields.set(qn("w:val"), "true")
    doc.settings.element.append(element_update_fields)


# ============================================================================
# 3. 封面、页眉页脚与原生自动目录 (TOC) 生成器
# ============================================================================
def setup_page_layout_and_header_footer(doc):
    """配置 A4 页面边距、页眉与页脚"""
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

        # 页眉
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(HEADER_TEXT)
        set_run_font(hrun, size_pt=8.5, color=COLOR_TEXT_MUTED)

        # 页脚
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Confidential & Proprietary · 莫伯 (MoBo) · 本地优先与风险隔离架构")
        set_run_font(frun, size_pt=8.5, color=COLOR_TEXT_MUTED)


def build_cover_page(doc):
    """构建机构级工程总册封面页与原生 Word 自动目录页"""
    for _ in range(3):
        doc.add_paragraph()

    p_tag = doc.add_paragraph()
    p_tag.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_tag = p_tag.add_run(DOC_TAG)
    set_run_font(r_tag, ascii_font="Consolas", size_pt=10, bold=True, color=COLOR_GOLD_ACCENT)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run(DOC_TITLE)
    set_run_font(r_title, size_pt=26, bold=True, color=COLOR_NAVY_DARK)
    add_bottom_border_to_paragraph(p_title, color_hex=HEX_GOLD_LINE, sz="18")

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(8)
    p_sub.paragraph_format.space_after = Pt(28)
    r_sub = p_sub.add_run(DOC_SUBTITLE)
    set_run_font(r_sub, size_pt=13, bold=True, color=COLOR_NAVY_MED)

    # 封面核心元数据卡片表
    meta_rows = COVER_META_ROWS or [
        ("文档编号", "MOBO-DUAL-STACK-MASTER-V2.1 (定稿版)"),
        ("核心计算底座", "阿里云 2核4G ECS (Ubuntu 24.04 LTS) + 4GB Swap + 6 容器硬限流 (常驻约 1.46GB)"),
        ("数据库隔离内核", "PostgreSQL 16 (pgvector) 单实例双库隔离：db_mobobrain & db_archive"),
        ("存储与灾备防线", "同地域双 OSS 桶内网免流 + 5 条冷热沉降规则 + Restic 每日 03:00 加密冷备"),
        ("自动化与终端引擎", "n8n“1+3”四端同源合成发布工作流 + mobo_cli.py v2.1 + macro_iv_collector.py"),
        ("多智能体协同", "Claude Code CLI (架构/DB/运维) + Codex CLI (脚本/安全审计) + Cursor (双站前端)"),
        ("生成日期", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]

    tbl = doc.add_table(rows=len(meta_rows), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color_hex=HEX_BORDER_GRAY, sz="4")

    for idx, (k, v) in enumerate(meta_rows):
        row = tbl.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Cm(4.2)
        c1.width = Cm(12.4)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF" if idx % 2 == 0 else HEX_ZEBRA_ROW)
        set_cell_margins(c0, 100, 100, 140, 140)
        set_cell_margins(c1, 100, 100, 140, 140)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        set_run_font(r0, size_pt=9.5, bold=True, color=COLOR_NAVY_DARK)

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        set_run_font(r1, size_pt=9.5, color=COLOR_TEXT_MAIN)

    doc.add_page_break()

    # 插入 Word 原生目录页 (TOC Field)
    p_toc_title = doc.add_paragraph()
    r_toc_title = p_toc_title.add_run("目  录 (Table of Contents)")
    set_run_font(r_toc_title, size_pt=16, bold=True, color=COLOR_NAVY_DARK)
    add_bottom_border_to_paragraph(p_toc_title, color_hex=HEX_GOLD_LINE, sz="12")

    p_toc_hint = doc.add_paragraph()
    r_hint = p_toc_hint.add_run("提示：在 Word 或 WPS 中打开本文档后，右键点击下方目录区域选择“更新域 (Update Field)”即可自动刷新完整章节页码。")
    set_run_font(r_hint, size_pt=9, italic=True, color=COLOR_TEXT_MUTED)

    p_toc = doc.add_paragraph()
    run_toc = p_toc.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")
    r_element = run_toc._r
    r_element.append(fldChar1)
    r_element.append(instrText)
    r_element.append(fldChar2)
    r_element.append(fldChar3)

    doc.add_page_break()


# ============================================================================
# 4. 富文本行内解析器 (支持 **加粗** 与 `行内代码`)
# ============================================================================
INLINE_TOKEN_RE = re.compile(r"(\*\*.*?\*\*|`[^`]+`)")


def append_formatted_inline_text(paragraph, text: str, base_size_pt=10.5, base_color=COLOR_TEXT_MAIN):
    """解析单行中的 **bold** 与 `inline_code` 并赋予对应字体样式"""
    tokens = INLINE_TOKEN_RE.split(text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = paragraph.add_run(tok[2:-2])
            set_run_font(r, size_pt=base_size_pt, bold=True, color=COLOR_NAVY_DARK)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1])
            set_run_font(
                r,
                ascii_font="Consolas",
                east_asia_font="Microsoft YaHei",
                size_pt=base_size_pt - 0.5,
                bold=True,
                color=COLOR_NAVY_MED,
            )
        else:
            r = paragraph.add_run(tok)
            set_run_font(r, size_pt=base_size_pt, color=base_color)


# ============================================================================
# 5. 代码块容器、表格与引用块渲染器
# ============================================================================
def add_code_block_box(doc, code_lines: list, lang_label: str = "code"):
    """将代码块渲染为带语言角标与浅灰底色的单格表格容器"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color_hex=HEX_BORDER_GRAY, sz="6")

    cell = tbl.rows[0].cells[0]
    cell.width = Cm(16.6)
    set_cell_background(cell, HEX_CODE_BG)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15

    if lang_label:
        r_lang = p.add_run(f"[{lang_label.upper()}]\n")
        set_run_font(r_lang, ascii_font="Consolas", size_pt=8.5, bold=True, color=COLOR_GOLD_ACCENT)

    code_text = "\n".join(code_lines).rstrip()
    r_code = p.add_run(code_text)
    set_run_font(
        r_code,
        ascii_font="Consolas",
        east_asia_font="Microsoft YaHei",
        size_pt=8.5,
        color=COLOR_CODE_TEXT,
    )

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(4)


def add_markdown_table(doc, table_lines: list):
    """解析 Markdown 表格行并渲染为华尔街深海蓝表头 + 斑马纹表格"""
    parsed_rows = []
    for line in table_lines:
        stripped = line.strip().strip("|")
        cells = [c.strip() for c in stripped.split("|")]
        # 跳过 | :--- | :--- | 分隔行
        if all(re.match(r"^:?-{2,}:?$", c.replace(" ", "")) for c in cells if c):
            continue
        parsed_rows.append(cells)

    if not parsed_rows:
        return

    num_cols = max(len(r) for r in parsed_rows)
    tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color_hex=HEX_BORDER_GRAY, sz="4")

    for r_idx, row_data in enumerate(parsed_rows):
        row = tbl.rows[r_idx]
        is_header = (r_idx == 0)
        bg_hex = HEX_NAVY_HEADER if is_header else ("FFFFFF" if r_idx % 2 == 1 else HEX_ZEBRA_ROW)

        for c_idx in range(num_cols):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_background(cell, bg_hex)
            set_cell_margins(cell, top=100, bottom=100, left=130, right=130)

            raw_val = row_data[c_idx] if c_idx < len(row_data) else ""
            raw_val = raw_val.replace("<br>", "\n").replace("<br/>", "\n")

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.2

            if is_header:
                clean_header = re.sub(r"\*\*|`", "", raw_val)
                r = p.add_run(clean_header)
                set_run_font(r, size_pt=9.5, bold=True, color=RGBColor(255, 255, 255))
            else:
                append_formatted_inline_text(p, raw_val, base_size_pt=9.0, base_color=COLOR_TEXT_MAIN)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(4)


def add_blockquote_box(doc, quote_lines: list):
    """将 > 引用块渲染为暖色提示框"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color_hex=HEX_GOLD_LINE, sz="6")

    cell = tbl.rows[0].cells[0]
    cell.width = Cm(16.6)
    set_cell_background(cell, HEX_QUOTE_BG)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)

    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.25
    combined = "\n".join(quote_lines)
    append_formatted_inline_text(p, combined, base_size_pt=9.5, base_color=COLOR_NAVY_DARK)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


# ============================================================================
# 6. Markdown 全文状态机解析器
# ============================================================================
def render_markdown_to_doc(doc, md_text: str):
    """逐行状态机解析 Markdown 文本并转换为 Word 样式组件"""
    lines = md_text.splitlines()
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 1. 跳过水平分割线 ---
        if re.match(r"^-{3,}$", stripped):
            i += 1
            continue

        # 2. 解析围栏代码块 ```lang ... ```
        if stripped.startswith("```"):
            lang = stripped[3:].strip() or "code"
            code_buf = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_buf.append(lines[i])
                i += 1
            i += 1  # 跳过闭合 ```
            add_code_block_box(doc, code_buf, lang_label=lang)
            continue

        # 3. 解析 Markdown 表格 | ... |
        if stripped.startswith("|") and stripped.endswith("|"):
            tbl_buf = []
            while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                tbl_buf.append(lines[i])
                i += 1
            add_markdown_table(doc, tbl_buf)
            continue

        # 4. 解析引用块 > ...
        if stripped.startswith(">"):
            q_buf = []
            while i < n and lines[i].strip().startswith(">"):
                q_buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            add_blockquote_box(doc, q_buf)
            continue

        # 5. 解析各级标题 (# / ## / ### / ####)
        heading_match = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            title_text = re.sub(r"\*\*|`", "", heading_match.group(2).strip())

            # 文档最顶部的总标题若已在封面呈现，跳过重复的一级大标题
            if level == 1 and "莫伯双站系统" in title_text and i < 5:
                i += 1
                continue

            h_para = doc.add_heading(level=min(level, 4))
            h_para.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            h_para.paragraph_format.space_after = Pt(6)
            h_run = h_para.add_run(title_text)

            if level == 1:
                set_run_font(h_run, size_pt=16, bold=True, color=COLOR_NAVY_DARK)
                add_bottom_border_to_paragraph(h_para, color_hex=HEX_GOLD_LINE, sz="12")
            elif level == 2:
                set_run_font(h_run, size_pt=14, bold=True, color=COLOR_NAVY_DARK)
                add_bottom_border_to_paragraph(h_para, color_hex=HEX_BORDER_GRAY, sz="6")
            elif level == 3:
                set_run_font(h_run, size_pt=12, bold=True, color=COLOR_NAVY_MED)
            else:
                set_run_font(h_run, size_pt=10.5, bold=True, color=COLOR_NAVY_MED)

            i += 1
            continue

        # 6. 解析无序列表 (- / * ) 与有序列表 (1. )
        bullet_match = re.match(r"^(\s*)([-\*\•]|\d+\.)\s+(.*)$", line)
        if bullet_match:
            indent_spaces = len(bullet_match.group(1))
            marker = bullet_match.group(2)
            item_text = bullet_match.group(3).strip()

            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6 if indent_spaces == 0 else 1.2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25

            prefix_str = "• " if marker in ("-", "*", "•") else f"{marker} "
            r_pre = p.add_run(prefix_str)
            set_run_font(r_pre, size_pt=10.0, bold=True, color=COLOR_GOLD_ACCENT)
            append_formatted_inline_text(p, item_text, base_size_pt=10.0)
            i += 1
            continue

        # 7. 空行处理
        if not stripped:
            i += 1
            continue

        # 8. 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.35
        append_formatted_inline_text(p, stripped, base_size_pt=10.5)
        i += 1


# ============================================================================
# 7. 从 /opt/mobo-stack 实盘目录自动扫描并组装总册 Markdown (--from-stack 模式)
# ============================================================================
def assemble_markdown_from_stack_dir(stack_dir: Path) -> str:
    """当指定 --from-stack 时，自动读取 /opt/mobo-stack 下的所有核心配置文件与源码合成总册"""
    chunks = [
        "## 第一章：多智能体共享架构宪章 (CLAUDE.md) 与任务交接账本 (TASK_STATE.md)",
    ]

    file_manifest = [
        ("1.1 多智能体架构宪章 (`CLAUDE.md`)", stack_dir / "CLAUDE.md", "markdown"),
        ("1.2 16 步原子任务交接账本 (`TASK_STATE.md`)", stack_dir / "TASK_STATE.md", "markdown"),
        ("## 第二章：双站网关与六容器硬限流编排 (`Caddyfile` & `docker-compose.yml`)", None, None),
        ("2.1 双域名 HTTPS 与 BasicAuth 网关 (`caddy/Caddyfile`)", stack_dir / "caddy/Caddyfile", "caddyfile"),
        ("2.2 六容器硬限流编排 (`docker-compose.yml`)", stack_dir / "docker-compose.yml", "yaml"),
        ("## 第三章：PostgreSQL 16 (pgvector) 双库物理隔离与自动化触发器全集", None, None),
        ("3.1 双库隔离初始化 (`01-init-dual-db.sql`)", stack_dir / "init-scripts/01-init-dual-db.sql", "sql"),
        ("3.2 莫伯之家核心表 (`02-archive-schema.sql`)", stack_dir / "init-scripts/02-archive-schema.sql", "sql"),
        ("3.3 莫伯之家时间线触发器 (`03-archive-triggers.sql`)", stack_dir / "init-scripts/03-archive-triggers.sql", "sql"),
        ("3.4 莫伯资本核心表 (`04-mobocapital-schema.sql`)", stack_dir / "init-scripts/04-mobocapital-schema.sql", "sql"),
        ("3.5 莫伯资本 NAV 高水位与风控熔断触发器 (`05-mobocapital-nav-and-triggers.sql`)", stack_dir / "init-scripts/05-mobocapital-nav-and-triggers.sql", "sql"),
        ("3.6 莫伯资本量化终端表 (`06-mobocapital-terminal-schema.sql`)", stack_dir / "init-scripts/06-mobocapital-terminal-schema.sql", "sql"),
        ("## 第四章：双栖终端录入引擎与每日量化采集脚本", None, None),
        ("4.1 双站终端总控工具 (`scripts/mobo_cli.py`)", stack_dir / "scripts/mobo_cli.py", "python"),
        ("4.2 宏观先行指标与期权波动率采集器 (`scripts/macro_iv_collector.py`)", stack_dir / "scripts/macro_iv_collector.py", "python"),
        ("4.3 OSS 冷热沉降与防盗链守卫 (`scripts/oss_guard_setup.py`)", stack_dir / "scripts/oss_guard_setup.py", "python"),
    ]

    for title, fpath, lang in file_manifest:
        if fpath is None:
            chunks.append(f"\n{title}\n")
            continue
        chunks.append(f"\n### {title}\n")
        if fpath.exists():
            content = fpath.read_text(encoding="utf-8")
            if lang == "markdown":
                chunks.append(content)
            else:
                chunks.append(f"```{lang}\n{content}\n```")
        else:
            chunks.append(f"> 待落盘文件路径：`{fpath}`")

    return "\n".join(chunks)


# ============================================================================
# 8. 主程序入口
# ============================================================================
def main():
    # 封面/页眉文案覆盖（模块级变量被 build_cover_page / setup_page_layout 读取）
    global DOC_TAG, DOC_TITLE, DOC_SUBTITLE, HEADER_TEXT, COVER_META_ROWS

    parser = argparse.ArgumentParser(
        description="无需 Pandoc，直接将莫伯双站系统总册渲染为机构级 Word (.docx) 文档"
    )
    parser.add_argument(
        "-i", "--input-md",
        type=str,
        default="MoBo_Dual_Stack_Master_v2.1.md",
        help="输入的 Markdown 总册文件路径 (默认: MoBo_Dual_Stack_Master_v2.1.md)",
    )
    parser.add_argument(
        "-o", "--output-docx",
        type=str,
        default="《莫伯双站系统全栈架构与部署总册_v2.1》.docx",
        help="输出的 Word 文件路径 (默认: 《莫伯双站系统全栈架构与部署总册_v2.1》.docx)",
    )
    parser.add_argument(
        "--from-clipboard",
        action="store_true",
        help="直接从系统剪贴板读取刚刚复制的总册 Markdown 内容并生成 Word",
    )
    parser.add_argument(
        "--from-stack",
        type=str,
        nargs="?",
        const="/opt/mobo-stack",
        help="直接扫描 /opt/mobo-stack 目录下的实盘代码与配置文件自动组装生成 Word 总册",
    )
    parser.add_argument(
        "--doc-tag",
        type=str,
        default=DOC_TAG,
        help="封面左上角的小标签文字",
    )
    parser.add_argument(
        "--doc-title",
        type=str,
        default=DOC_TITLE,
        help="封面主标题",
    )
    parser.add_argument(
        "--doc-subtitle",
        type=str,
        default=DOC_SUBTITLE,
        help="封面副标题（可用 \\n 分行）",
    )
    parser.add_argument(
        "--header-text",
        type=str,
        default=HEADER_TEXT,
        help="页眉文字",
    )
    parser.add_argument(
        "--cover-meta",
        action="append",
        default=[],
        metavar="KEY=VALUE",
        help="封面元数据卡行，可重复；给了就替换默认卡片",
    )

    args = parser.parse_args()
    md_content = ""

    DOC_TAG = args.doc_tag
    DOC_TITLE = args.doc_title
    DOC_SUBTITLE = args.doc_subtitle.replace("\\n", "\n")
    HEADER_TEXT = args.header_text
    COVER_META_ROWS = []
    for item in args.cover_meta:
        if "=" in item:
            k, v = item.split("=", 1)
            COVER_META_ROWS.append((k.strip(), v.strip()))
        else:
            print(f"⚠️ 忽略格式不合法的 --cover-meta：{item}")

    if args.from_clipboard:
        try:
            import tkinter as tk
            root = tk.Tk()
            root.withdraw()
            md_content = root.clipboard_get()
            root.destroy()
            print(f"📋 已从系统剪贴板读取 {len(md_content)} 字符。")
        except Exception as e:
            print(f"❌ 从剪贴板读取失败 ({e})，请改用 -i MoBo_Dual_Stack_Master_v2.1.md 文件方式运行。")
            sys.exit(1)
    elif args.from_stack:
        stack_path = Path(args.from_stack).expanduser().resolve()
        print(f"🔍 正在扫描生产目录 {stack_path} 并自动汇编全栈源码...")
        md_content = assemble_markdown_from_stack_dir(stack_path)
    else:
        in_path = Path(args.input_md).expanduser().resolve()
        if not in_path.exists():
            print(
                f"⚠️ 未找到默认输入文件: {in_path}\n"
                f"👉 方式 1: 请先点击上一条回答右上角的【复制】按钮，然后直接运行:\n"
                f"   python build_master_docx.py --from-clipboard\n"
                f"👉 方式 2: 将上一条回答保存为 MoBo_Dual_Stack_Master_v2.1.md 后直接运行:\n"
                f"   python build_master_docx.py"
            )
            sys.exit(1)
        md_content = in_path.read_text(encoding="utf-8")
        print(f"📄 已加载 Markdown 源文件: {in_path} ({len(md_content)} 字符)")

    doc = Document()
    setup_page_layout_and_header_footer(doc)
    enable_auto_update_fields(doc)
    build_cover_page(doc)
    render_markdown_to_doc(doc, md_content)

    out_path = Path(args.output_docx).expanduser().resolve()
    doc.save(str(out_path))
    print(f"✅ 机构级 Word 总册生成完毕: {out_path}")


if __name__ == "__main__":
    main()